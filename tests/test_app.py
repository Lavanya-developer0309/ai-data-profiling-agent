"""Run: .venv/bin/python -m pytest -q  (or .venv/bin/python tests/test_app.py)"""
import json
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import auth
from app import create_app


SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def client(**config):
    db = tempfile.NamedTemporaryFile(suffix=".db", delete=False).name
    auth._failed.clear()
    return create_app({"DATABASE": db, "TESTING": True, "UPLOADS": tempfile.mkdtemp(), **config}).test_client()


def csrf(c, url="/login"):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(url).text).group(1)


def register(c, **kw):
    data = dict(mode="register", full_name="Lavanya Natarajan", username="lavanya",
                password="Profile2026", confirm="Profile2026", csrf=csrf(c))
    data.update(kw)
    return c.post("/login", data=data)


def test_main_requires_login():
    c = client()
    r = c.get("/app")
    assert r.status_code == 302 and "/login" in r.location


def test_register_then_main_page():
    c = client()
    r = register(c)
    assert r.status_code == 302 and r.location.endswith("/app")
    page = c.get("/app").text
    assert "Hi Lavanya" in page and "Drag and drop" in page
    assert c.get("/api/status").json["ai"]["provider"]


def test_register_validation():
    c = client()
    r = register(c, username="x", password="short", confirm="short")
    assert r.status_code == 400
    assert "3-30 characters" in r.text and "At least 8 characters" in r.text
    r = register(c, confirm="Different1")
    assert "Passwords do not match" in r.text
    register(c)
    c.post("/logout", data={"csrf": csrf(c, "/app")})
    assert "already taken" in register(c).text


def test_password_is_hashed():
    c = client()
    register(c)
    with c.application.app_context():
        row = auth.db().execute("SELECT password_hash FROM users").fetchone()
    assert "Profile2026" not in row[0] and row[0].startswith(("scrypt:", "pbkdf2:"))


def test_login_logout_and_lockout():
    c = client()
    register(c)
    c.post("/logout", data={"csrf": csrf(c, "/app")})
    assert c.get("/app").status_code == 302

    ok = c.post("/login", data=dict(mode="signin", username="LAVANYA", password="Profile2026", csrf=csrf(c)))
    assert ok.location.endswith("/app")
    c.post("/logout", data={"csrf": csrf(c, "/app")})

    for _ in range(5):
        r = c.post("/login", data=dict(mode="signin", username="lavanya", password="wrong", csrf=csrf(c)))
        assert r.status_code == 400
    r = c.post("/login", data=dict(mode="signin", username="lavanya", password="Profile2026", csrf=csrf(c)))
    assert "Too many failed attempts" in r.text


def test_csrf_required():
    c = client()
    r = c.post("/login", data=dict(mode="signin", username="a", password="b"))
    assert r.status_code == 400 and "Form expired" in r.text


def test_open_redirect_blocked():
    c = client()
    c.get("/login")
    data = dict(mode="register", full_name="Test User", username="tester", password="Profile2026",
                confirm="Profile2026", csrf=csrf(c))
    r = c.post("/login?next=//evil.example", data=data)
    assert r.location.endswith("/app")


# ---------------------------------------------------------------- step 2

def upload(c, name="SRC_MEMBER.txt", as_name=None):
    with open(SAMPLES / name, "rb") as f:
        return c.post("/api/runs", data={"file": (f, as_name or name)}, headers={"X-CSRF-Token": csrf(c, "/app")},
                      content_type="multipart/form-data")


def test_propose_without_snowflake_uses_builtin_rules():
    c = client(ASK_AI="off")
    register(c)
    r = upload(c)
    assert r.status_code == 200, r.json
    data = r.json
    assert data["rows"] == 3938 and len(data["rules"]) == 18
    assert data["info"]["source"] == "rules"
    rules = {x["column"]: x for x in data["rules"]}
    assert rules["FIRST_NAME"]["max_len"] == 15 and rules["LAST_NAME"]["max_len"] == 35
    assert rules["PHONE"]["type"] == "PHONE" and rules["DATE_OF_BIRTH"]["type"] == "BIRTH_DATE"


def test_propose_with_ai_and_ssn_masked():
    seen = {}

    def fake_ai(prompt):
        seen["prompt"] = prompt
        return 'Here you go:\n```json\n[{"column": "MEMBER_ID", "type": "ID", "required": true, "unique": true,' \
               ' "reason": "Primary key"}, {"column": "FIRST_NAME", "type": "NAME", "max_len": 15, "reason": "First name"},' \
               ' {"column": "GENDER_CODE", "type": "code", "allowed_values": ["M", "F", "U"]}, {"column": "PHONE", "type": "BOGUS"}]\n```'
    c = client(ASK_AI=fake_ai)
    register(c)
    data = upload(c).json
    rules = {x["column"]: x for x in data["rules"]}
    assert data["info"]["source"] == "ai" and "Built-in rules used for" in data["info"]["message"]
    assert rules["MEMBER_ID"]["source"] == "ai" and rules["MEMBER_ID"]["unique"]
    assert rules["GENDER_CODE"]["type"] == "CODE" and rules["GENDER_CODE"]["allowed_values"] == ["M", "F", "U"]
    assert rules["PHONE"]["type"] == "TEXT"            # unknown type from AI is made safe
    assert rules["LAST_NAME"]["source"] == "rules"     # AI skipped it -> built-in rule
    ssn_line = next(l for l in seen["prompt"].splitlines() if l.startswith("SSN:"))
    assert not re.search(r"[0-8]", ssn_line), ssn_line  # real digits never sent


def test_guardrails_fix_typical_ai_mistakes():
    def sloppy_ai(prompt):  # the mistakes qwen3 actually made on SRC_MEMBER.txt
        return json.dumps([
            {"column": "MEMBER_ID", "type": "ID", "required": True, "min_len": 10, "max_len": 12, "unique": True},
            {"column": "SUBSCRIBER_ID", "type": "ID", "unique": True},
            {"column": "GENDER_CODE", "type": "CODE", "allowed_values": ["M", "F", "1", "2", "X"]},
            {"column": "PLAN_CODE", "type": "CODE", "allowed_values": ["P200", "P300", "P100", "D100", "M100", "M200", "P999"]},
            {"column": "STATE", "type": "STATE", "allowed_values": ["TX", "CO"]},
            {"column": "DATE_OF_BIRTH", "type": "DATE", "min_len": 10, "max_len": 10},
            {"column": "FIRST_NAME", "type": "NAME", "max_len": 15},
        ])
    c = client(ASK_AI=sloppy_ai)
    register(c)
    data = upload(c).json
    rules = {x["column"]: x for x in data["rules"]}
    assert rules["MEMBER_ID"]["min_len"] is None and rules["MEMBER_ID"]["max_len"] == 12  # IDs are 8 chars: min 10 wrong, max 12 fine
    assert rules["MEMBER_ID"]["unique"] and not rules["SUBSCRIBER_ID"]["unique"]
    assert rules["GENDER_CODE"]["allowed_values"] == ["M", "F", "U"]
    assert rules["PLAN_CODE"]["allowed_values"] == [] and "copied" in rules["PLAN_CODE"]["reason"]
    assert rules["STATE"]["allowed_values"] == []
    assert rules["DATE_OF_BIRTH"]["type"] == "BIRTH_DATE" and rules["DATE_OF_BIRTH"]["max_len"] is None
    assert rules["DATE_OF_BIRTH"]["date_format"] == "MM/DD/YYYY"
    assert rules["FIRST_NAME"]["max_len"] == 15 and "Adjusted" not in rules["FIRST_NAME"]["reason"]
    assert "Guardrails corrected" in data["info"]["message"]


def test_ai_failure_falls_back():
    def broken_ai(prompt):
        raise RuntimeError("Cortex model not available in region")
    c = client(ASK_AI=broken_ai)
    register(c)
    data = upload(c).json
    assert data["info"]["source"] == "rules" and "Cortex model not available" in data["info"]["message"]
    assert len(data["rules"]) == 18


def test_approve_rules_and_validation():
    c = client(ASK_AI="off")
    register(c)
    data = upload(c, "SRC_PROVIDER.txt").json
    token = csrf(c, "/app")
    rules = data["rules"]
    rules[3].update(max_len=20, source="user")
    ok = c.put(f"/api/runs/{data['run_id']}/rules", json={"rules": rules}, headers={"X-CSRF-Token": token})
    assert ok.status_code == 200 and ok.json["rules"][3]["max_len"] == 20

    rules[3].update(min_len=30)
    bad = c.put(f"/api/runs/{data['run_id']}/rules", json={"rules": rules}, headers={"X-CSRF-Token": token})
    assert bad.status_code == 400 and "min length is bigger" in bad.json["error"]
    missing = c.put(f"/api/runs/{data['run_id']}/rules", json={"rules": rules[:-1]}, headers={"X-CSRF-Token": token})
    assert missing.status_code == 400 and "No rule for" in missing.json["error"]
    assert c.put(f"/api/runs/{data['run_id']}/rules", json={"rules": rules}).status_code == 400  # no csrf


def test_runs_are_private_and_bad_files_rejected():
    c = client(ASK_AI="off")
    register(c)
    run_id = upload(c).json["run_id"]
    c.post("/logout", data={"csrf": csrf(c, "/app")})
    register(c, username="someone", full_name="Someone Else")
    r = c.put(f"/api/runs/{run_id}/rules", json={"rules": []}, headers={"X-CSRF-Token": csrf(c, "/app")})
    assert r.status_code == 404
    assert c.put("/api/runs/../../etc/rules", json={}, headers={"X-CSRF-Token": csrf(c, "/app")}).status_code == 404
    import io
    r = c.post("/api/runs", data={"file": (io.BytesIO(b"x"), "evil.exe")}, headers={"X-CSRF-Token": csrf(c, "/app")},
               content_type="multipart/form-data")
    assert r.status_code == 400


# ---------------------------------------------------------------- step 3

def test_execute_requires_approved_rules():
    c = client(ASK_AI="off")
    register(c)
    run_id = upload(c).json["run_id"]
    r = c.post(f"/api/runs/{run_id}/execute", headers={"X-CSRF-Token": csrf(c, "/app")})
    assert r.status_code == 400 and "Approve the checks first" in r.json["error"]


def test_execute_in_snowflake_finds_planted_issues():
    """Real Snowflake run - counts must match the issues planted in the sample file."""
    import os
    if not os.getenv("SNOWFLAKE_ACCOUNT"):
        print("  (skipped - Snowflake not set up)")
        return
    c = client(ASK_AI="off")
    register(c)
    data = upload(c, as_name="ZZ_TEST_SRC_MEMBER.txt").json   # own name, so real SRC_MEMBER results are not replaced
    token = csrf(c, "/app")
    assert c.put(f"/api/runs/{data['run_id']}/rules", json={"rules": data["rules"]}, headers={"X-CSRF-Token": token}).status_code == 200
    r = c.post(f"/api/runs/{data['run_id']}/execute", headers={"X-CSRF-Token": token})
    assert r.status_code == 200, r.json
    found = {(x["column"], x["check"]): x["failed"] for x in r.json["checks"]}
    assert found[("GENDER_CODE", "Not in: M, F, U")] == 16          # 12 numeric + 4 X
    assert found[("MEMBER_ID", "Duplicates")] == 10
    assert found[("DATE_OF_BIRTH", "Missing values")] == 5
    assert found[("DATE_OF_BIRTH", "Invalid date / not MM/DD/YYYY")] == 3
    assert found[("FIRST_NAME", "Special characters")] == 8
    assert found[("LAST_NAME", "Leading / trailing spaces")] == 3
    assert found[("MIDDLE_INITIAL", "Leading / trailing spaces")] == 0  # blanks are nulls, not spaces
    ssn = [x for x in r.json["checks"] if x["column"] == "SSN"]
    assert all(not any(ch.isdigit() for ch in e[:-4]) for x in ssn for e in x["examples"])  # masked
    assert r.json["rows"] == 3938
    saved = r.json["saved"]
    assert saved["ok"], saved
    assert saved["rows_saved"]["PROFILE_FAILED_ROW"] == sum(x["failed"] for x in r.json["checks"])  # < 1000 each here
    assert saved["export_queries"][0][0] == "Issues for this file" and "V_PROFILE_RESULTS" in saved["export_queries"][0][1]

    # drill-down: the exact rows behind the counts
    gender = next(x for x in r.json["checks"] if x["check"] == "Not in: M, F, U")
    d = c.get(f"/api/runs/{data['run_id']}/checks/{gender['id']}/rows").json
    assert d["table"] == "ZZ_TEST_SRC_MEMBER" and d["column"] == "GENDER_CODE" and d["key_column"] == "MEMBER_ID"
    assert len(d["rows"]) == 16 and {x["value"] for x in d["rows"]} == {"1", "2", "X"}
    assert all(x["length"] == 1 and x["key"].startswith("M") for x in d["rows"])
    from profiling.rules import read_source
    src, _ = read_source(SAMPLES / "SRC_MEMBER.txt")
    assert all(src.iloc[x["row"] - 1]["GENDER_CODE"] == x["value"] for x in d["rows"])   # row numbers point at the right rows
    assert all(not any(ch.isdigit() for ch in x["record"]["SSN"][:-4]) for x in d["rows"])  # SSN masked in full record
    dup = next(x for x in r.json["checks"] if x["check"] == "Duplicates")
    d = c.get(f"/api/runs/{data['run_id']}/checks/{dup['id']}/rows").json
    assert len(d["rows"]) == 10

    # permanent tables hold the same rows; then remove this test's rows
    import sf
    from profiling.store import SCHEMA, TABLES
    conn = sf.connect()
    try:
        cur = conn.cursor()
        cur.execute(f"SELECT KEY_VALUE, COLUMN_VALUE FROM {SCHEMA}.V_PROFILE_FAILED_ROWS "
                    "WHERE SOURCE_FILE = 'ZZ_TEST_SRC_MEMBER.txt' AND COLUMN_NAME = 'GENDER_CODE' ORDER BY ROW_NO")
        assert [tuple(x) for x in cur.fetchall()] == [(x["key"], x["value"]) for x in c.get(
            f"/api/runs/{data['run_id']}/checks/{gender['id']}/rows").json["rows"]]
        for t in TABLES:
            cur.execute(f"DELETE FROM {SCHEMA}.{t} WHERE SOURCE_FILE = 'ZZ_TEST_SRC_MEMBER.txt'")
    finally:
        conn.close()


# ---------------------------------------------------------------- step 4

def _checked_run():
    """A finished run with realistic results, without needing Snowflake."""
    from profiling.rules import guess_rules, read_source
    df, _ = read_source(SAMPLES / "SRC_MEMBER.txt")
    checks = [
        {"id": 0, "column": "GENDER_CODE", "type": "CODE", "check": "Not in: M, F, U", "failed": 16, "pct": 0.41, "examples": ["1", "2", "X"]},
        {"column": "MEMBER_ID", "type": "ID", "check": "Duplicates", "failed": 10, "pct": 0.25, "examples": ["M0001927"]},
        {"column": "SSN", "type": "SSN", "check": "Dashes or other characters", "failed": 9, "pct": 0.23, "examples": ["XXX-XX-2346"]},
        {"column": "LAST_NAME", "type": "NAME", "check": "Leading / trailing spaces", "failed": 3, "pct": 0.08, "examples": ['"  Harris "']},
        {"column": "FIRST_NAME", "type": "NAME", "check": "Contains digits", "failed": 0, "pct": 0, "examples": []},
    ]
    results = {"rows": 3938, "checks_run": 5, "checks_failed": 4, "issues": 38, "columns_with_issues": 4,
               "seconds": 4.0, "load_seconds": 2.0, "checks": checks,
               "profile": [{"column": c, "type": "TEXT", "rows": 3938, "nulls": 0, "null_pct": 0, "distinct": 5,
                            "min_len": 1, "max_len": 9, "top_values": [{"value": "A", "count": 3}]} for c in df.columns]}
    return {"filename": "SRC_MEMBER.txt", "rules": guess_rules(df), "results": results, "checked_at": "2026-10-21T12:00:00"}


def test_report_severity_and_ssn_examples_hidden_from_ai():
    from profiling import report
    facts = report.build_facts(_checked_run())
    assert [f["severity"] for f in facts["findings"]] == ["High", "High", "Medium", "Low"]  # dashes can be cleansed -> Medium
    ssn = next(f for f in facts["findings"] if f["column"] == "SSN")
    assert ssn["examples"] == []


def test_report_rejects_made_up_numbers():
    from profiling import report
    calls = []

    def lying_ai(prompt):
        calls.append(prompt)
        return {"headline": "About 40 issues found", "overview": "Gender has 17 bad rows (0.5%).",
                "findings": [{"id": 1, "explanation": "16 rows bad", "recommendation": "Fix them"}], "next_steps": ["Fix"]}
    summary, info = report.write_summary(_checked_run(), lying_ai)
    assert info["source"] == "standard" and "17" in info["message"] and "40" in info["message"]
    assert len(calls) == 2 and "not in the results" in calls[1]          # retried once with feedback
    assert "38 data issues" in summary["headline"]


def test_report_accepts_correct_numbers_on_retry():
    from profiling import report
    answers = iter([
        {"headline": "39 issues", "overview": "x", "findings": [], "next_steps": []},
        {"headline": "3,938 rows checked, 38 issues found", "overview": "Gender: 16 rows (0.41%).",
         "findings": [{"id": 1, "explanation": "16 rows use codes 1, 2 and X.", "recommendation": "Map them."}],
         "next_steps": ["Agree a crosswalk"]},
    ])
    summary, info = report.write_summary(_checked_run(), lambda p: next(answers))
    assert info["source"] == "ai" and "second attempt" in info["message"]
    assert summary["findings"][0]["explanation"] == "16 rows use codes 1, 2 and X."
    assert summary["findings"][1]["explanation"].startswith("10 rows in MEMBER_ID")   # skipped by AI -> standard text


def test_excel_report():
    import io
    from openpyxl import load_workbook
    from profiling import report
    run = _checked_run()
    summary = report.fallback_summary(report.build_facts(run))
    wb = load_workbook(io.BytesIO(report.build_excel(run, summary, "Lavanya Natarajan")))
    assert wb.sheetnames == ["Summary", "All checks", "Column profile", "Rules used"]
    ws = wb["Summary"]
    assert ws["A1"].value == "Data Profiling Report" and "Lavanya Natarajan" in ws["A2"].value
    assert ws["A5"].value == 3938 and ws["E5"].value == 38
    assert wb["All checks"].max_row == 6 and wb["All checks"]["D2"].value == 16


def test_summary_and_download_endpoints():
    c = client(ASK_AI="off", ASK_SUMMARY_AI="off")
    register(c)
    data = upload(c).json
    token = csrf(c, "/app")
    assert c.post(f"/api/runs/{data['run_id']}/summary", headers={"X-CSRF-Token": token}).status_code == 400  # not checked yet
    path = Path(c.application.config["UPLOADS"]) / data["run_id"] / "run.json"
    run = json.loads(path.read_text())
    run.update(_checked_run(), run_id=data["run_id"], user_id=run["user_id"], status="checked")
    path.write_text(json.dumps(run))
    r = c.post(f"/api/runs/{data['run_id']}/summary", headers={"X-CSRF-Token": token})
    assert r.status_code == 200 and r.json["info"]["source"] == "standard" and len(r.json["findings"]) == 4
    assert c.get(f"/api/runs/{data['run_id']}/checks/0/rows").status_code == 409   # no failures.json for this run
    x = c.get(f"/api/runs/{data['run_id']}/report.xlsx")
    assert x.status_code == 200 and x.data[:2] == b"PK" and "SRC_MEMBER_profiling_report_" in x.headers["Content-Disposition"]


# ---------------------------------------------------------------- pipeline page

def _approvals_db(rows):
    import sqlite3
    path = tempfile.NamedTemporaryFile(suffix=".db", delete=False).name
    with sqlite3.connect(path) as conn:
        conn.execute("""CREATE TABLE approvals (id INTEGER PRIMARY KEY, run_id TEXT, batch TEXT, step INTEGER, tool TEXT,
                        reason TEXT, status TEXT DEFAULT 'PENDING', note TEXT, requested_at TEXT, answered_at TEXT)""")
        conn.executemany("INSERT INTO approvals (run_id, batch, step, tool, reason, requested_at) VALUES (?,?,?,?,?,?)", rows)
    return path


def test_pipeline_approvals_flow():
    db = _approvals_db([("run1", "BATCH_009", 6, "dbt_build", "profiling looks fine", "2026-10-21T20:00:00"),
                        ("run1", "BATCH_009", 7, "load_target", "dbt passed", "2026-10-21T20:01:00")])
    c = client(ASK_AI="off", APPROVALS_DB=db)
    register(c)
    assert c.get("/pipeline").status_code == 200 and "Waiting for your approval" in c.get("/pipeline").text
    data = c.get("/api/pipeline/approvals").json
    assert [a["tool"] for a in data["pending"]] == ["dbt_build", "load_target"]
    token = csrf(c, "/pipeline")
    first, second = data["pending"][0]["id"], data["pending"][1]["id"]
    assert c.post(f"/api/pipeline/approvals/{first}", json={"approve": True}, headers={"X-CSRF-Token": token}).json["ok"]
    r = c.post(f"/api/pipeline/approvals/{second}", json={"approve": False, "note": "client resending"}, headers={"X-CSRF-Token": token})
    assert r.json["ok"]
    assert c.post(f"/api/pipeline/approvals/{first}", json={"approve": False}, headers={"X-CSRF-Token": token}).status_code == 409
    assert c.post(f"/api/pipeline/approvals/{first}", json={"approve": True}).status_code == 400          # no csrf
    recent = {a["tool"]: a for a in c.get("/api/pipeline/approvals").json["recent"]}
    assert recent["dbt_build"]["status"] == "APPROVED" and recent["load_target"]["status"] == "REJECTED"
    assert recent["load_target"]["note"] == "client resending (by lavanya)"


def test_pipeline_requires_login_and_missing_db_is_empty():
    c = client(ASK_AI="off", APPROVALS_DB="/nonexistent/approvals.db")
    assert c.get("/pipeline").status_code == 302
    register(c)
    assert c.get("/api/pipeline/approvals").json == {"pending": [], "recent": []}


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn(); print("PASS", name)
